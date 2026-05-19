import streamlit as st
import pandas as pd
import re
import uuid
import unicodedata
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ImportError:
    Document = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import gspread
    from google.oauth2.service_account import Credentials
except ImportError:
    gspread = None
    Credentials = None


# ============================================================
# CONFIGURACIÓN GENERAL
# ============================================================

st.set_page_config(
    page_title="LexMonitor AI",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

CSV_LOCAL = "expedientes_lexmonitor.csv"

HEADERS = [
    "codigo",
    "fecha_creacion",
    "estado",
    "demandante",
    "demandado",
    "cuantia",
    "concepto_deuda",
    "documentacion_detectada",
    "categoria_art_812",
    "respuesta_deudor",
    "fecha_respuesta",
    "motivo_oposicion",
    "accion_recomendada",
    "borrador"
]


# ============================================================
# ESTILOS
# ============================================================

st.markdown(
    """
    <style>
    [data-testid="stSidebar"] {
        display: none;
    }

    [data-testid="collapsedControl"] {
        display: none;
    }

    .block-container {
        padding-top: 1.5rem;
        padding-left: 3rem;
        padding-right: 3rem;
    }

    .main {
        background: linear-gradient(135deg, #f8fafc 0%, #eef2ff 100%);
    }

    .topbar {
        display: flex;
        justify-content: space-between;
        align-items: center;
        margin-bottom: 1.5rem;
    }

    .topbar-title {
        font-size: 1.1rem;
        font-weight: 800;
        color: #111827;
    }

    .topbar-user {
        font-size: 0.9rem;
        color: #6b7280;
    }

    .login-wrapper {
        max-width: 520px;
        margin: 5rem auto 2rem auto;
        padding: 2.8rem;
        background: white;
        border-radius: 30px;
        box-shadow: 0 25px 70px rgba(15, 23, 42, 0.15);
        border: 1px solid #e5e7eb;
        text-align: center;
    }

    .login-icon {
        width: 82px;
        height: 82px;
        margin: 0 auto 1rem auto;
        border-radius: 50%;
        background: #1e3a8a;
        color: white;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 38px;
        font-weight: bold;
    }

    .login-badge {
        display: inline-block;
        background: #eff6ff;
        color: #1d4ed8;
        padding: 0.4rem 0.9rem;
        border-radius: 999px;
        font-size: 0.82rem;
        font-weight: 700;
        margin-bottom: 1rem;
    }

    .login-title {
        font-size: 2rem;
        font-weight: 850;
        color: #111827;
        margin-bottom: 0.4rem;
    }

    .login-subtitle {
        font-size: 1rem;
        color: #6b7280;
        margin-bottom: 1.8rem;
    }

    .blank-selection {
        min-height: 72vh;
        display: flex;
        align-items: center;
        justify-content: center;
    }

    .selection-card {
        background: white;
        border-radius: 30px;
        padding: 3rem;
        width: 100%;
        max-width: 900px;
        box-shadow: 0 25px 70px rgba(15, 23, 42, 0.12);
        border: 1px solid #e5e7eb;
    }

    .selection-title {
        text-align: center;
        font-size: 1.7rem;
        font-weight: 850;
        color: #111827;
        margin-bottom: 0.4rem;
    }

    .selection-subtitle {
        text-align: center;
        color: #6b7280;
        margin-bottom: 2rem;
    }

    .role-card {
        padding: 2rem;
        border-radius: 24px;
        background: #ffffff;
        border: 1px solid #e5e7eb;
        box-shadow: 0 12px 35px rgba(15, 23, 42, 0.08);
        text-align: center;
        min-height: 230px;
        margin-bottom: 1rem;
    }

    .role-card-debt {
        padding: 2rem;
        border-radius: 24px;
        background: #fff7ed;
        border: 2px solid #f97316;
        box-shadow: 0 18px 45px rgba(249, 115, 22, 0.22);
        text-align: center;
        min-height: 230px;
        margin-bottom: 1rem;
    }

    .role-icon {
        font-size: 2.4rem;
        margin-bottom: 0.8rem;
    }

    .role-title {
        font-size: 1.45rem;
        font-weight: 850;
        color: #111827;
        margin-bottom: 0.5rem;
    }

    .role-text {
        font-size: 0.98rem;
        color: #6b7280;
        line-height: 1.5;
    }

    .debt-alert {
        background: #fff7ed;
        border: 1px solid #fdba74;
        border-left: 6px solid #f97316;
        border-radius: 18px;
        padding: 1rem 1.3rem;
        margin-bottom: 1.5rem;
        color: #7c2d12;
        font-weight: 650;
    }

    .panel {
        background: white;
        border-radius: 22px;
        padding: 1.6rem;
        border: 1px solid #e5e7eb;
        box-shadow: 0 12px 30px rgba(15, 23, 42, 0.07);
        margin-bottom: 1.2rem;
    }

    .status-box {
        background: #f8fafc;
        border-left: 5px solid #1d4ed8;
        padding: 1rem 1.2rem;
        border-radius: 14px;
        margin-bottom: 1rem;
    }

    .success-box {
        background: #ecfdf5;
        border-left: 5px solid #059669;
        padding: 1rem 1.2rem;
        border-radius: 14px;
        margin-bottom: 1rem;
    }

    .warning-box {
        background: #fffbeb;
        border-left: 5px solid #d97706;
        padding: 1rem 1.2rem;
        border-radius: 14px;
        margin-bottom: 1rem;
    }

    div.stButton > button {
        border-radius: 13px;
        font-weight: 750;
        padding: 0.72rem 1rem;
    }

    div[data-testid="stFileUploader"] {
        background: #f9fafb;
        border: 1px dashed #cbd5e1;
        border-radius: 16px;
        padding: 1rem;
    }
    </style>
    """,
    unsafe_allow_html=True
)


# ============================================================
# ESTADO DE SESIÓN
# ============================================================

if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if "perfil" not in st.session_state:
    st.session_state.perfil = None

if "usuario_nombre" not in st.session_state:
    st.session_state.usuario_nombre = ""

if "registro_demandado" not in st.session_state:
    st.session_state.registro_demandado = None

if "mostrar_oposicion" not in st.session_state:
    st.session_state.mostrar_oposicion = False


# ============================================================
# UTILIDADES
# ============================================================

def normalizar_texto(texto):
    texto = str(texto or "").strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto


def google_sheets_disponible():
    return (
        gspread is not None
        and Credentials is not None
        and "SPREADSHEET_ID" in st.secrets
        and "gcp_service_account" in st.secrets
    )


def get_worksheet():
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive.readonly"
    ]

    credentials = Credentials.from_service_account_info(
        st.secrets["gcp_service_account"],
        scopes=scopes
    )

    client = gspread.authorize(credentials)
    spreadsheet = client.open_by_key(st.secrets["SPREADSHEET_ID"])

    try:
        worksheet = spreadsheet.worksheet("expedientes")
    except gspread.WorksheetNotFound:
        worksheet = spreadsheet.add_worksheet(
            title="expedientes",
            rows=1000,
            cols=len(HEADERS)
        )
        worksheet.append_row(HEADERS)

    values = worksheet.get_all_values()
    if not values:
        worksheet.append_row(HEADERS)

    return worksheet


def asegurar_csv_local():
    path = Path(CSV_LOCAL)
    if not path.exists():
        pd.DataFrame(columns=HEADERS).to_csv(path, index=False)


def cargar_registros():
    if google_sheets_disponible():
        try:
            ws = get_worksheet()
            records = ws.get_all_records()
            return pd.DataFrame(records, columns=HEADERS)
        except Exception as e:
            st.warning(f"No se pudo conectar con Google Sheets. Usando CSV local. Error: {e}")

    asegurar_csv_local()
    return pd.read_csv(CSV_LOCAL, dtype=str).fillna("")


def guardar_registro(registro):
    registro_completo = {h: str(registro.get(h, "")) for h in HEADERS}

    if google_sheets_disponible():
        try:
            ws = get_worksheet()
            ws.append_row([registro_completo[h] for h in HEADERS])
            return "google_sheets"
        except Exception as e:
            st.warning(f"No se pudo guardar en Google Sheets. Se guardará en CSV local. Error: {e}")

    asegurar_csv_local()
    df = pd.read_csv(CSV_LOCAL, dtype=str).fillna("")
    df = pd.concat([df, pd.DataFrame([registro_completo])], ignore_index=True)
    df.to_csv(CSV_LOCAL, index=False)
    return "csv_local"


def actualizar_registro(codigo, cambios):
    df = cargar_registros()

    if df.empty or "codigo" not in df.columns:
        return False

    mask = df["codigo"].astype(str).str.upper() == codigo.upper()

    if not mask.any():
        return False

    for clave, valor in cambios.items():
        if clave in df.columns:
            df.loc[mask, clave] = str(valor)

    if google_sheets_disponible():
        try:
            ws = get_worksheet()
            values = df[HEADERS].fillna("").values.tolist()
            ws.clear()
            ws.append_row(HEADERS)
            if values:
                ws.append_rows(values)
            return True
        except Exception as e:
            st.warning(f"No se pudo actualizar Google Sheets. Se actualizará CSV local. Error: {e}")

    asegurar_csv_local()
    df.to_csv(CSV_LOCAL, index=False)
    return True


def buscar_por_codigo(codigo):
    df = cargar_registros()

    if df.empty or "codigo" not in df.columns:
        return None

    mask = df["codigo"].astype(str).str.upper() == codigo.upper()

    if not mask.any():
        return None

    return df[mask].iloc[0].to_dict()


def buscar_deudas_por_nombre(nombre):
    df = cargar_registros()

    if df.empty or "demandado" not in df.columns:
        return pd.DataFrame(columns=HEADERS)

    nombre_norm = normalizar_texto(nombre)

    estados_pendientes = [
        "pendiente de respuesta del deudor",
        "pendiente de subsanación",
        "sin pago ni oposición"
    ]

    df["demandado_norm"] = df["demandado"].apply(normalizar_texto)
    df["estado_norm"] = df["estado"].apply(normalizar_texto)

    coincidencias = df[
        (df["demandado_norm"] == nombre_norm)
        & (df["estado_norm"].isin(estados_pendientes))
    ].copy()

    coincidencias = coincidencias.drop(columns=["demandado_norm", "estado_norm"], errors="ignore")

    return coincidencias


# ============================================================
# LECTURA DE DOCUMENTOS
# ============================================================

def leer_txt(archivo):
    return archivo.read().decode("utf-8", errors="ignore")


def leer_docx(archivo):
    if Document is None:
        return ""

    doc = Document(archivo)
    textos = []

    for p in doc.paragraphs:
        if p.text.strip():
            textos.append(p.text.strip())

    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [celda.text.strip() for celda in fila.cells]
            textos.append(" | ".join(celdas))

    return "\n".join(textos)


def leer_pdf(archivo):
    if PyPDF2 is None:
        return ""

    lector = PyPDF2.PdfReader(archivo)
    textos = []

    for pagina in lector.pages:
        texto = pagina.extract_text()
        if texto:
            textos.append(texto)

    return "\n".join(textos)


def leer_documento(archivo):
    nombre = archivo.name.lower()

    if nombre.endswith(".txt"):
        return leer_txt(archivo)

    if nombre.endswith(".docx"):
        return leer_docx(archivo)

    if nombre.endswith(".pdf"):
        return leer_pdf(archivo)

    return ""


# ============================================================
# EXTRACCIÓN DE DATOS
# ============================================================

def limpiar_numero(texto_numero):
    if not texto_numero:
        return 0.0

    texto_numero = str(texto_numero).lower()
    texto_numero = texto_numero.replace("euros", "")
    texto_numero = texto_numero.replace("euro", "")
    texto_numero = texto_numero.replace("€", "")
    texto_numero = texto_numero.strip()

    if "," in texto_numero and "." in texto_numero:
        if texto_numero.find(",") < texto_numero.find("."):
            texto_numero = texto_numero.replace(",", "")
        else:
            texto_numero = texto_numero.replace(".", "").replace(",", ".")
    elif "," in texto_numero:
        partes = texto_numero.split(",")
        if len(partes[-1]) == 2:
            texto_numero = texto_numero.replace(",", ".")
        else:
            texto_numero = texto_numero.replace(",", "")
    elif "." in texto_numero:
        partes = texto_numero.split(".")
        if len(partes[-1]) != 2:
            texto_numero = texto_numero.replace(".", "")

    try:
        return float(texto_numero)
    except ValueError:
        return 0.0


def buscar_patron(texto, patrones):
    for patron in patrones:
        match = re.search(patron, texto, re.IGNORECASE | re.MULTILINE)
        if match:
            return match.group(1).strip()
    return ""


def detectar_documentos_art_812(texto):
    texto = texto.lower()

    categorias = {
        "Documento firmado, sellado o validado por el deudor": [
            "firmado por el deudor",
            "firma del deudor",
            "sello del deudor",
            "firma electrónica",
            "señal electrónica",
            "documento firmado"
        ],
        "Facturas, albaranes, certificaciones, telegramas, fax u otros documentos habituales": [
            "factura",
            "facturas",
            "albarán",
            "albaranes",
            "certificación",
            "certificaciones",
            "telegrama",
            "telegramas",
            "fax",
            "recibo",
            "recibos",
            "estado de cuenta",
            "comprobante"
        ],
        "Documentos comerciales que acreditan relación anterior duradera": [
            "relación comercial anterior",
            "relación anterior duradera",
            "relación contractual continuada",
            "contrato marco",
            "contrato de suministro",
            "historial de pedidos",
            "documentos comerciales",
            "relación mercantil"
        ],
        "Certificación de impago de comunidad de propietarios": [
            "certificación de impago",
            "gastos comunes",
            "comunidad de propietarios",
            "cuotas comunitarias"
        ],
    }

    detectados = {}

    for categoria, palabras in categorias.items():
        coincidencias = [p for p in palabras if p in texto]
        detectados[categoria] = coincidencias

    hay_documento_valido = any(len(v) > 0 for v in detectados.values())

    return hay_documento_valido, detectados


def categorias_detectadas_texto(documentos_art_812):
    categorias = []

    for categoria, coincidencias in documentos_art_812.items():
        if coincidencias:
            categorias.append(categoria)

    return "; ".join(categorias)


def extraer_datos_demanda(texto_demanda, texto_documentos):
    texto_total = f"{texto_demanda}\n\n{texto_documentos}"
    texto_unido = re.sub(r"\s+", " ", texto_total)

    demandante = buscar_patron(
        texto_unido,
        [
            r"demandante[:\s]+(.+?)(?:demandado|deudor|contra|frente a|,|\.)",
            r"acreedor[:\s]+(.+?)(?:demandado|deudor|contra|frente a|,|\.)",
            r"promovido por[:\s]+(.+?)(?:contra|frente a|,|\.)",
            r"a instancia de[:\s]+(.+?)(?:contra|frente a|,|\.)",
        ]
    )

    demandado = buscar_patron(
        texto_unido,
        [
            r"demandado[:\s]+(.+?)(?:,|\.)",
            r"deudor[:\s]+(.+?)(?:,|\.)",
            r"contra[:\s]+(.+?)(?:,|\.)",
            r"frente a[:\s]+(.+?)(?:,|\.)",
        ]
    )

    cuantia_txt = buscar_patron(
        texto_unido,
        [
            r"cuantía[:\s]+(?:de\s*)?(?:€\s*)?([\d\.,]+)",
            r"importe[:\s]+(?:de\s*)?(?:€\s*)?([\d\.,]+)",
            r"cantidad[:\s]+(?:de\s*)?(?:€\s*)?([\d\.,]+)",
            r"reclama(?:\s+la)?\s+cantidad\s+de\s+(?:€\s*)?([\d\.,]+)",
            r"por\s+importe\s+de\s+(?:€\s*)?([\d\.,]+)",
        ]
    )

    cuantia = limpiar_numero(cuantia_txt)

    concepto_deuda = buscar_patron(
        texto_unido,
        [
            r"concepto(?:\s+de\s+la\s+deuda)?[:\s]+(.{10,250})",
            r"la deuda deriva de[:\s]+(.{10,250})",
            r"deuda derivada de[:\s]+(.{10,250})",
            r"por los siguientes hechos[:\s]+(.{10,250})",
        ]
    )

    hay_documento_deuda, documentos_art_812 = detectar_documentos_art_812(texto_documentos)

    datos_faltantes = []

    if not demandante:
        datos_faltantes.append("Demandante / acreedor")
    if not demandado:
        datos_faltantes.append("Demandado / deudor")
    if cuantia <= 0:
        datos_faltantes.append("Cuantía")
    if not hay_documento_deuda:
        datos_faltantes.append("Documentación acreditativa art. 812 LEC")

    cumple_requisitos_auto = (
        bool(demandante)
        and bool(demandado)
        and cuantia > 0
        and hay_documento_deuda
    )

    return {
        "demandante": demandante,
        "demandado": demandado,
        "cuantia": cuantia,
        "concepto_deuda": concepto_deuda,
        "hay_documento_deuda": hay_documento_deuda,
        "documentos_art_812": documentos_art_812,
        "categoria_art_812": categorias_detectadas_texto(documentos_art_812),
        "datos_faltantes": datos_faltantes,
        "cumple_requisitos_auto": cumple_requisitos_auto,
    }


# ============================================================
# BORRADORES
# ============================================================

def generar_codigo():
    year = datetime.now().year
    sufijo = uuid.uuid4().hex[:6].upper()
    return f"MON-{year}-{sufijo}"


def generar_borrador_admision(registro):
    return f"""
AUTO DE ADMISIÓN DE SOLICITUD MONITORIA

Código de expediente: {registro['codigo']}

Visto el escrito presentado por {registro['demandante']} frente a {registro['demandado']}, por importe de {registro['cuantia']} euros, y examinada la documentación aportada, se aprecia inicialmente que la solicitud contiene los datos básicos necesarios y documentación acreditativa de la deuda.

Documentación detectada:
{registro['categoria_art_812']}

En consecuencia, procede admitir la solicitud monitoria y requerir al deudor para que pague la cantidad reclamada o formule oposición en el plazo legalmente previsto.

Documento generado automáticamente para revisión humana.
""".strip()


def generar_borrador_subsanacion(registro):
    return f"""
AUTO DE REQUERIMIENTO DE SUBSANACIÓN

Código de expediente: {registro['codigo']}

Examinada la solicitud presentada por {registro['demandante']} frente a {registro['demandado']}, se aprecia que no constan todos los elementos necesarios para su admisión inicial.

Actuación recomendada:
{registro['accion_recomendada']}

Documento generado automáticamente para revisión humana.
""".strip()


def generar_borrador_pago(registro):
    return f"""
DILIGENCIA DE PAGO Y ARCHIVO

Código de expediente: {registro['codigo']}

Constando que la parte demandada ha seleccionado la opción de pago respecto de la reclamación formulada por {registro['demandante']}, procede tener por atendida la reclamación y acordar el archivo de las actuaciones, previa comprobación del pago.

Documento generado automáticamente para revisión humana.
""".strip()


def generar_borrador_oposicion(registro, motivo):
    return f"""
REGISTRO DE OPOSICIÓN

Código de expediente: {registro['codigo']}

La parte demandada manifiesta su oposición al pago reclamado por {registro['demandante']}.

Cantidad reclamada:
{registro['cuantia']} euros

Motivo de oposición:
{motivo}

Actuación recomendada:
Dar traslado de la oposición y continuar por el cauce procesal correspondiente según la cuantía.

Documento generado automáticamente para revisión humana.
""".strip()


def generar_borrador_no_comparece(registro):
    return f"""
AUTO DE EJECUCIÓN POR FALTA DE PAGO U OPOSICIÓN

Código de expediente: {registro['codigo']}

Constando que la parte demandada no ha pagado ni ha formulado oposición dentro del escenario simulado, procede continuar con la actuación correspondiente por falta de comparecencia.

Actuación recomendada:
Iniciar ejecución por vía de apremio, previa comprobación del vencimiento del plazo legal.

Documento generado automáticamente para revisión humana.
""".strip()


# ============================================================
# LOGIN
# ============================================================

def pantalla_login():
    st.markdown(
        """
        <div class="login-wrapper">
            <div class="login-icon">⚖️</div>
            <div class="login-badge">Acceso seguro</div>
            <div class="login-title">LexMonitor AI</div>
            <div class="login-subtitle">
                Identificación mediante certificado digital
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

    col1, col2, col3 = st.columns([1, 2, 1])

    with col2:
        nombre = st.text_input(
            "Nombre y apellidos",
            placeholder="Introduce tu nombre completo"
        )

        certificado = st.file_uploader(
            "Certificado digital",
            type=["pdf", "txt", "cer", "crt", "p12", "pfx"]
        )

        entrar = st.button(
            "Acceder",
            use_container_width=True,
            disabled=not nombre.strip() or certificado is None
        )

        if entrar:
            st.session_state.autenticado = True
            st.session_state.usuario_nombre = nombre.strip()
            st.session_state.perfil = None
            st.rerun()


# ============================================================
# CABECERA SIN MENÚ LATERAL
# ============================================================

def topbar():
    col1, col2 = st.columns([8, 1.4])

    with col1:
        st.markdown(
            f"""
            <div class="topbar">
                <div>
                    <div class="topbar-title">LexMonitor AI</div>
                    <div class="topbar-user">Sesión iniciada como: {st.session_state.usuario_nombre}</div>
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

    with col2:
        if st.button("Cerrar sesión", use_container_width=True):
            st.session_state.autenticado = False
            st.session_state.perfil = None
            st.session_state.usuario_nombre = ""
            st.session_state.registro_demandado = None
            st.session_state.mostrar_oposicion = False
            st.rerun()


# ============================================================
# PANTALLA EN BLANCO CON DEMANDANTE / DEMANDADO
# ============================================================

def pantalla_perfiles():
    deudas = buscar_deudas_por_nombre(st.session_state.usuario_nombre)
    tiene_deudas = not deudas.empty

    st.markdown('<div class="blank-selection">', unsafe_allow_html=True)
    st.markdown('<div class="selection-card">', unsafe_allow_html=True)

    st.markdown(
        """
        <div class="selection-title">¿Cómo quieres acceder?</div>
        <div class="selection-subtitle">Selecciona el perfil con el que vas a operar.</div>
        """,
        unsafe_allow_html=True
    )

    if tiene_deudas:
        codigos = ", ".join(deudas["codigo"].astype(str).tolist())

        st.markdown(
            f"""
            <div class="debt-alert">
                Tienes una deuda pendiente asociada a tu nombre.
                <br>
                Código de deuda: <b>{codigos}</b>
            </div>
            """,
            unsafe_allow_html=True
        )

    col1, col2 = st.columns(2)

    with col1:
        st.markdown(
            """
            <div class="role-card">
                <div class="role-icon">📄</div>
                <div class="role-title">Demandante / Acreedor</div>
                <div class="role-text">
                    Presentar una solicitud monitoria y aportar documentación.
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button("Entrar como Demandante", use_container_width=True):
            st.session_state.perfil = "demandante"
            st.rerun()

    with col2:
        clase = "role-card-debt" if tiene_deudas else "role-card"

        texto = (
            "Tienes una reclamación pendiente. Consulta el expediente y selecciona una actuación."
            if tiene_deudas
            else "Consultar una reclamación mediante código y responder."
        )

        st.markdown(
            f"""
            <div class="{clase}">
                <div class="role-icon">💼</div>
                <div class="role-title">Demandado / Deudor</div>
                <div class="role-text">
                    {texto}
                </div>
            </div>
            """,
            unsafe_allow_html=True
        )

        if st.button("Entrar como Demandado", use_container_width=True):
            st.session_state.perfil = "demandado"

            if tiene_deudas:
                primer_registro = deudas.iloc[0].to_dict()
                st.session_state.registro_demandado = primer_registro

            st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)


# ============================================================
# DEMANDANTE
# ============================================================

def pantalla_demandante():
    topbar()

    if st.button("← Volver", use_container_width=False):
        st.session_state.perfil = None
        st.rerun()

    st.title("Zona del Demandante / Acreedor")

    modo = st.radio(
        "Elige cómo quieres presentar la solicitud",
        [
            "Subir PDF/DOCX/TXT de plantilla rellena",
            "Rellenar formulario desde la app"
        ],
        horizontal=True
    )

    texto_demanda = ""
    texto_documentos = ""

    datos_manual = {
        "demandante": st.session_state.usuario_nombre,
        "demandado": "",
        "cuantia": 0.0,
        "concepto_deuda": ""
    }

    if modo == "Subir PDF/DOCX/TXT de plantilla rellena":
        st.subheader("1. Subir solicitud monitoria")

        demanda_file = st.file_uploader(
            "Sube la demanda o solicitud principal",
            type=["pdf", "docx", "txt"],
            key="demanda_principal"
        )

        if demanda_file is not None:
            texto_demanda = leer_documento(demanda_file)

            with st.expander("Ver texto leído de la solicitud"):
                st.text_area("Texto detectado", texto_demanda, height=220)

    else:
        st.subheader("1. Rellenar formulario de solicitud monitoria")

        c1, c2 = st.columns(2)

        datos_manual["demandante"] = c1.text_input(
            "Demandante / acreedor",
            value=st.session_state.usuario_nombre
        )

        datos_manual["demandado"] = c2.text_input(
            "Demandado / deudor"
        )

        datos_manual["cuantia"] = st.number_input(
            "Cuantía reclamada",
            min_value=0.0,
            step=100.0
        )

        datos_manual["concepto_deuda"] = st.text_area(
            "Concepto de la deuda",
            placeholder="Ejemplo: deuda derivada de factura impagada por prestación de servicios..."
        )

        texto_demanda = f"""
        Demandante: {datos_manual['demandante']}.
        Demandado: {datos_manual['demandado']}.
        Cuantía: {datos_manual['cuantia']}.
        Concepto de la deuda: {datos_manual['concepto_deuda']}.
        Solicito la tramitación de proceso monitorio.
        """

    st.subheader("2. Subir documentos acreditativos")

    documentos_files = st.file_uploader(
        "Sube facturas, albaranes, certificaciones, documentos firmados o documentos comerciales",
        type=["pdf", "docx", "txt"],
        accept_multiple_files=True,
        key="documentos_acreditativos"
    )

    if documentos_files:
        textos_docs = []

        for doc_file in documentos_files:
            texto_doc = leer_documento(doc_file)
            textos_docs.append(
                f"\n\n--- DOCUMENTO ACREDITATIVO: {doc_file.name} ---\n{texto_doc}"
            )

        texto_documentos = "\n".join(textos_docs)

        with st.expander("Ver texto leído de los documentos"):
            st.text_area("Documentos detectados", texto_documentos, height=220)

    st.subheader("3. Comprobación y generación de código")

    if modo == "Rellenar formulario desde la app":
        datos_extraidos = {
            "demandante": datos_manual["demandante"],
            "demandado": datos_manual["demandado"],
            "cuantia": datos_manual["cuantia"],
            "concepto_deuda": datos_manual["concepto_deuda"],
        }

        hay_doc, documentos_art_812 = detectar_documentos_art_812(texto_documentos)

        datos_extraidos["hay_documento_deuda"] = hay_doc
        datos_extraidos["documentos_art_812"] = documentos_art_812
        datos_extraidos["categoria_art_812"] = categorias_detectadas_texto(documentos_art_812)

        datos_faltantes = []
        if not datos_extraidos["demandante"]:
            datos_faltantes.append("Demandante / acreedor")
        if not datos_extraidos["demandado"]:
            datos_faltantes.append("Demandado / deudor")
        if datos_extraidos["cuantia"] <= 0:
            datos_faltantes.append("Cuantía")
        if not hay_doc:
            datos_faltantes.append("Documentación acreditativa art. 812 LEC")

        datos_extraidos["datos_faltantes"] = datos_faltantes
        datos_extraidos["cumple_requisitos_auto"] = len(datos_faltantes) == 0

    else:
        datos_extraidos = extraer_datos_demanda(texto_demanda, texto_documentos)

    c1, c2, c3 = st.columns(3)

    c1.metric("Demandante", datos_extraidos["demandante"] or "No detectado")
    c2.metric("Demandado", datos_extraidos["demandado"] or "No detectado")
    c3.metric("Cuantía", f"{float(datos_extraidos['cuantia']):,.2f} €")

    if datos_extraidos["hay_documento_deuda"]:
        st.markdown(
            """
            <div class="success-box">
                Documentación acreditativa detectada.
            </div>
            """,
            unsafe_allow_html=True
        )
    else:
        st.markdown(
            """
            <div class="warning-box">
                No se ha detectado documentación acreditativa suficiente.
            </div>
            """,
            unsafe_allow_html=True
        )

    with st.expander("Ver categorías documentales detectadas"):
        docs_812 = datos_extraidos.get("documentos_art_812", {})
        for categoria, coincidencias in docs_812.items():
            if coincidencias:
                st.success(f"{categoria}: {', '.join(coincidencias)}")
            else:
                st.info(f"{categoria}: no detectado")

    if datos_extraidos["datos_faltantes"]:
        st.warning("Datos pendientes: " + ", ".join(datos_extraidos["datos_faltantes"]))

    generar = st.button(
        "Generar código de demanda",
        use_container_width=True,
        disabled=not bool(datos_extraidos["demandante"]) or not bool(datos_extraidos["demandado"])
    )

    if generar:
        codigo = generar_codigo()
        cumple = datos_extraidos["cumple_requisitos_auto"]

        if cumple:
            estado = "Pendiente de respuesta del deudor"
            accion = "Requerir al deudor para pagar u oponerse"
            documentacion_detectada = "Sí"
        else:
            estado = "Pendiente de subsanación"
            accion = "Requerir subsanación documental"
            documentacion_detectada = "No"

        registro = {
            "codigo": codigo,
            "fecha_creacion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "estado": estado,
            "demandante": datos_extraidos["demandante"],
            "demandado": datos_extraidos["demandado"],
            "cuantia": str(datos_extraidos["cuantia"]),
            "concepto_deuda": datos_extraidos.get("concepto_deuda", ""),
            "documentacion_detectada": documentacion_detectada,
            "categoria_art_812": datos_extraidos.get("categoria_art_812", ""),
            "respuesta_deudor": "",
            "fecha_respuesta": "",
            "motivo_oposicion": "",
            "accion_recomendada": accion,
            "borrador": ""
        }

        if cumple:
            borrador = generar_borrador_admision(registro)
        else:
            borrador = generar_borrador_subsanacion(registro)

        registro["borrador"] = borrador

        destino = guardar_registro(registro)

        st.success(f"Código generado correctamente: {codigo}")

        if destino == "google_sheets":
            st.info("Registro guardado en Google Sheets.")
        else:
            st.info("Registro guardado en CSV local.")

        st.subheader("Borrador generado")
        st.text_area("Borrador", borrador, height=330)

        st.download_button(
            "Descargar borrador",
            data=borrador,
            file_name=f"borrador_{codigo}.txt",
            mime="text/plain"
        )


# ============================================================
# DEMANDADO
# ============================================================

def pantalla_demandado():
    topbar()

    if st.button("← Volver", use_container_width=False):
        st.session_state.perfil = None
        st.session_state.registro_demandado = None
        st.rerun()

    st.title("Zona del Demandado / Deudor")

    registro_precargado = st.session_state.get("registro_demandado")

    if registro_precargado:
        st.markdown(
            f"""
            <div class="debt-alert">
                Se ha detectado una deuda pendiente asociada a tu nombre.
                <br>
                Código de deuda: <b>{registro_precargado.get('codigo', '')}</b>
            </div>
            """,
            unsafe_allow_html=True
        )

    codigo_default = registro_precargado.get("codigo", "") if registro_precargado else ""

    codigo = st.text_input(
        "Código de demanda",
        value=codigo_default,
        placeholder="Ejemplo: MON-2026-ABC123"
    )

    buscar = st.button("Buscar reclamación", use_container_width=True)

    if buscar and codigo:
        registro = buscar_por_codigo(codigo.strip())

        if registro is None:
            st.error("No se ha encontrado ninguna reclamación con ese código.")
            return

        st.session_state.registro_demandado = registro

    registro = st.session_state.get("registro_demandado")

    if not registro:
        return

    st.subheader("Resumen de la reclamación")

    c1, c2, c3 = st.columns(3)
    c1.metric("Código", registro.get("codigo", ""))
    c2.metric("Estado", registro.get("estado", ""))
    c3.metric("Cantidad reclamada", f"{registro.get('cuantia', '')} €")

    st.markdown(
        f"""
        <div class="panel">
            <b>Demandante / acreedor:</b> {registro.get('demandante', '')}<br>
            <b>Demandado / deudor:</b> {registro.get('demandado', '')}<br>
            <b>Concepto de la deuda:</b> {registro.get('concepto_deuda', '')}<br>
            <b>Documentación aportada:</b> {registro.get('categoria_art_812', '')}<br>
            <b>Actuación recomendada actual:</b> {registro.get('accion_recomendada', '')}
        </div>
        """,
        unsafe_allow_html=True
    )

    st.subheader("Selecciona una actuación")

    col1, col2, col3 = st.columns(3)

    with col1:
        if st.button("Pagar deuda", use_container_width=True):
            borrador = generar_borrador_pago(registro)

            actualizar_registro(
                registro["codigo"],
                {
                    "estado": "Pagado",
                    "respuesta_deudor": "Pagar deuda",
                    "fecha_respuesta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "accion_recomendada": "Archivar expediente tras comprobar pago",
                    "borrador": borrador
                }
            )

            st.success("Respuesta registrada: pago de deuda.")
            st.text_area("Borrador generado", borrador, height=300)

    with col2:
        if st.button("Oponerse al pago", use_container_width=True):
            st.session_state.mostrar_oposicion = True

    with col3:
        if st.button("No pagar ni comparecer", use_container_width=True):
            borrador = generar_borrador_no_comparece(registro)

            actualizar_registro(
                registro["codigo"],
                {
                    "estado": "Sin pago ni oposición",
                    "respuesta_deudor": "No pagar ni comparecer",
                    "fecha_respuesta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "accion_recomendada": "Iniciar ejecución por vía de apremio tras comprobar plazo",
                    "borrador": borrador
                }
            )

            st.warning("Respuesta registrada: no pago ni comparecencia.")
            st.text_area("Borrador generado", borrador, height=300)

    if st.session_state.get("mostrar_oposicion", False):
        st.subheader("Formulario de oposición")

        with st.form("form_oposicion"):
            motivo = st.text_area(
                "Motivo de oposición",
                placeholder="Explica brevemente por qué te opones al pago..."
            )

            reconoce_parte = st.checkbox("Reconozco parte de la deuda")

            cantidad_reconocida = 0.0
            if reconoce_parte:
                cantidad_reconocida = st.number_input(
                    "Cantidad reconocida",
                    min_value=0.0,
                    step=100.0
                )

            enviar = st.form_submit_button("Registrar oposición")

        if enviar:
            motivo_final = motivo

            if reconoce_parte:
                motivo_final += f"\nCantidad reconocida: {cantidad_reconocida}"

            borrador = generar_borrador_oposicion(registro, motivo_final)

            actualizar_registro(
                registro["codigo"],
                {
                    "estado": "Oposición presentada",
                    "respuesta_deudor": "Oponerse al pago",
                    "fecha_respuesta": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "motivo_oposicion": motivo_final,
                    "accion_recomendada": "Dar traslado de la oposición y continuar según cuantía",
                    "borrador": borrador
                }
            )

            st.success("Oposición registrada correctamente.")
            st.text_area("Borrador generado", borrador, height=320)


# ============================================================
# EJECUCIÓN PRINCIPAL
# ============================================================

if not st.session_state.autenticado:
    pantalla_login()
    st.stop()

if st.session_state.perfil is None:
    topbar()
    pantalla_perfiles()
    st.stop()

if st.session_state.perfil == "demandante":
    pantalla_demandante()

elif st.session_state.perfil == "demandado":
    pantalla_demandado()
